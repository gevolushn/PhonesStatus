"""
xlsx_parser.py — ручний режим таблиці: читання телефонів з локального .xlsx.

Контракт (спільний із `parsers/sheets_parser.py`): `dict[номер, 'ON'|'OFF']`.

Що змінилось у 1.2.0: аркуш і колонки більше НЕ вшиті в код — приходять параметрами
з конфігу. Причина конкретна: розкладка таблиці вже переїжджала (до 2026-08 номер
лежав у колонці B, потім B стала «Поверх», номер поїхав у F), і кожен такий переїзд
означав правку коду й реліз. Дефолти лишились ті самі — F і H.

Автовизначення аркуша збережене як ДЕФОЛТ (порожній `sheet`), але тепер його можна
перекрити явно. Це не косметика: на робочому файлі евристика фактично не спрацьовує —
єдиний кандидат зветься «Copy of Телефони 27.08.2026», а «copy of» стоїть у списку
виключень, тож усі кандидати відпадають і рятує лише фолбек на активний аркуш.
"""
from __future__ import annotations

import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from core.logger import log

# Дефолтна розкладка — та сама, що була вшита до 1.2.0.
DEFAULT_COL_NUMBER: str = "F"   # «Внутрішній»
DEFAULT_COL_STATUS: str = "H"   # «Статус»

# Аркуш обирається за назвою: має містити «телефон» і не містити жодного з виключень.
_SHEET_KEYWORD: str = "телефон"
_SHEET_EXCLUDE: tuple[str, ...] = ("старий", "copy of", "експорт")

# Внутрішній номер — ЦІЛКОМ числовий, 3-5 цифр.
_RE_EXTENSION = re.compile(r"^\d{3,5}$")

# Приймаються лише ці два статуси (після .strip().upper()).
_VALID_STATUSES: frozenset[str] = frozenset({"ON", "OFF"})


def _col_index(letter: str, fallback: str) -> int:
    """Літера колонки Excel → 0-based індекс у кортежі рядка. Хибне значення → дефолт."""
    try:
        return column_index_from_string(str(letter).strip().upper()) - 1
    except Exception:
        log.warning(f"Некоректна літера колонки {letter!r} — беру дефолт {fallback}.")
        return column_index_from_string(fallback) - 1


def detect_sheet(sheet_names: list[str]) -> str:
    """
    Автовизначення аркуша за назвою. Повертає "" якщо кандидатів немає.

    Кандидати — назви з «телефон», без слів із `_SHEET_EXCLUDE`. Якщо кандидатів
    кілька, береться ОСТАННІЙ у порядку аркушів (він найновіший).
    """
    candidates = [
        name for name in sheet_names
        if _SHEET_KEYWORD in name.lower()
        and not any(bad in name.lower() for bad in _SHEET_EXCLUDE)
    ]
    return candidates[-1] if candidates else ""


def list_sheets(path: str) -> list[str]:
    """Назви всіх аркушів книги — для випадайки в налаштуваннях."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def preview_columns(path: str, sheet: str = "", max_columns: int = 26) -> list[tuple[str, str]]:
    """
    Пари (літера, підказка) для випадайки вибору колонки: [('F', 'Внутрішній'), ...].

    Підказка — перша непорожня комірка колонки згори (типово заголовок). Якщо
    колонка порожня — підказка порожня, але літера все одно пропонується.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        hints: dict[int, str] = {}
        for row_index, row in enumerate(ws.iter_rows(values_only=True)):
            if row_index > 20:            # заголовок не буває нижче — далі лише дані
                break
            for col_index, value in enumerate(row[:max_columns]):
                if col_index not in hints and value not in (None, ""):
                    hints[col_index] = str(value).strip()
        return [
            (get_column_letter(i + 1), hints.get(i, ""))
            for i in range(max_columns)
        ]
    finally:
        wb.close()


def read_xlsx(
    path: str,
    sheet: str = "",
    col_number: str = DEFAULT_COL_NUMBER,
    col_status: str = DEFAULT_COL_STATUS,
) -> dict[str, str]:
    """
    Повертає `dict[внутрішній_номер, 'ON'|'OFF']`.

    `sheet` порожній → автовизначення (`detect_sheet`), далі фолбек на активний аркуш.
    Номер нормалізується через `int(float(...))`, щоб `1026.0` з Excel став `1026`.

    ⚠️ Комірка номера має бути ЧИСТИМ числом: `901 (Binotel)` не розпізнається — це
    номери зовнішньої АТС Binotel, а не внутрішні extension'и FreePBX, і вони свідомо
    поза порівнянням. Рядки зі статусом поза ON/OFF тихо пропускаються.
    """
    log.info(f"Читаю xlsx: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet and sheet in wb.sheetnames:
            ws = wb[sheet]
            log.info(f"Аркуш обрано явно: '{sheet}'")
        else:
            if sheet:
                log.warning(f"Аркуш '{sheet}' не знайдено — переходжу до автовизначення.")
            detected = detect_sheet(wb.sheetnames)
            if detected:
                ws = wb[detected]
                log.info(f"Аркуш визначено автоматично: '{detected}'")
            else:
                ws = wb.active
                log.warning(f"Підходящий аркуш не знайдено, беру активний: '{ws.title}'")

        idx_number = _col_index(col_number, DEFAULT_COL_NUMBER)
        idx_status = _col_index(col_status, DEFAULT_COL_STATUS)
        min_columns = max(idx_number, idx_status) + 1

        phones: dict[str, str] = {}
        for row in ws.iter_rows(values_only=True):
            if len(row) < min_columns:
                continue
            raw_number = row[idx_number]
            if raw_number is None:
                continue
            try:
                number = str(int(float(str(raw_number))))
            except (ValueError, TypeError):
                continue
            if not _RE_EXTENSION.match(number):
                continue

            raw_status = row[idx_status]
            status = str(raw_status).strip().upper() if raw_status else ""
            if status not in _VALID_STATUSES:
                continue

            phones[number] = status
    finally:
        wb.close()

    log.info(
        f"Зчитано {len(phones)} номерів з таблиці "
        f"(колонки {col_number}/{col_status})."
    )
    return phones
